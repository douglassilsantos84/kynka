"""
Importador de mapas de quantidades para demandas da Kynka.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from kynka.application.demand.demand_service import (
    DemandService,
)
from kynka.domain.inventory import (
    InventoryRepository,
)


@dataclass(slots=True)
class QuantityMapMissingMaterial:
    """
    Material presente no mapa, mas ausente no estoque.
    """

    row: int
    code: str
    name: str
    quantity: float
    unit: str


@dataclass(slots=True)
class QuantityMapImportResult:
    """
    Resultado da importação de um mapa de quantidades.
    """

    filename: str

    total_rows: int = 0
    imported: int = 0
    skipped: int = 0

    missing_materials: list[
        QuantityMapMissingMaterial
    ] = field(default_factory=list)

    errors: list[str] = field(
        default_factory=list
    )


class QuantityMapImporter:
    """
    Importa um mapa de quantidades XLSX para uma demanda.

    Colunas esperadas:

    Código | Material | Quantidade | Unidade
    """

    def __init__(
        self,
        demand_service: DemandService,
        inventory_repository: InventoryRepository,
    ) -> None:
        self._demand_service = demand_service
        self._inventory_repository = (
            inventory_repository
        )

    def import_file(
        self,
        demand_id: int,
        file_path: str | Path,
    ) -> QuantityMapImportResult:
        """
        Importa um arquivo XLSX para a demanda.
        """

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"Arquivo não encontrado: {path}"
            )

        if path.suffix.lower() != ".xlsx":
            raise ValueError(
                "O mapa de quantidades deve estar "
                "no formato .xlsx."
            )

        # Confirma que a demanda existe antes de
        # iniciar a leitura da planilha.
        self._demand_service.get_demand(
            demand_id
        )

        result = QuantityMapImportResult(
            filename=path.name
        )

        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        try:
            worksheet = workbook.active

            header_row = next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                None,
            )

            if header_row is None:
                raise ValueError(
                    "A planilha está vazia."
                )

            columns = self._resolve_columns(
                header_row
            )

            for row_number, values in enumerate(
                worksheet.iter_rows(
                    min_row=2,
                    values_only=True,
                ),
                start=2,
            ):
                if self._row_is_empty(values):
                    continue

                result.total_rows += 1

                try:
                    code = self._text(
                        self._cell(
                            values,
                            columns["code"],
                        )
                    )

                    name = self._text(
                        self._cell(
                            values,
                            columns["name"],
                        )
                    )

                    quantity_value = self._cell(
                        values,
                        columns["quantity"],
                    )

                    unit = self._text(
                        self._cell(
                            values,
                            columns["unit"],
                        )
                    )

                    if not code:
                        raise ValueError(
                            "Código do material vazio."
                        )

                    if not name:
                        raise ValueError(
                            "Nome do material vazio."
                        )

                    if not unit:
                        raise ValueError(
                            "Unidade vazia."
                        )

                    quantity = self._quantity(
                        quantity_value
                    )

                    if quantity <= 0:
                        raise ValueError(
                            "Quantidade deve ser maior "
                            "que zero."
                        )

                    material = (
                        self._inventory_repository
                        .get_by_code(code)
                    )

                    if material is None:
                        result.missing_materials.append(
                            QuantityMapMissingMaterial(
                                row=row_number,
                                code=code,
                                name=name,
                                quantity=quantity,
                                unit=unit,
                            )
                        )

                        result.skipped += 1
                        continue

                    self._demand_service.set_requirement(
                        demand_id=demand_id,
                        material_code=material.code,
                        required_quantity=quantity,
                    )

                    result.imported += 1

                except Exception as error:
                    result.skipped += 1

                    result.errors.append(
                        f"Linha {row_number}: {error}"
                    )

        finally:
            workbook.close()

        return result

    @staticmethod
    def _resolve_columns(
        headers,
    ) -> dict[str, int]:
        normalized = {}

        for index, header in enumerate(headers):
            key = QuantityMapImporter._normalize_header(
                header
            )

            if key:
                normalized[key] = index

        aliases = {
            "code": {
                "codigo",
                "cod",
                "code",
                "codigo material",
                "cod material",
            },
            "name": {
                "material",
                "nome",
                "descricao",
                "designacao",
                "produto",
            },
            "quantity": {
                "quantidade",
                "qtd",
                "qtde",
                "quantity",
            },
            "unit": {
                "unidade",
                "un",
                "unit",
            },
        }

        resolved = {}

        for field_name, possibilities in aliases.items():
            for alias in possibilities:
                if alias in normalized:
                    resolved[field_name] = (
                        normalized[alias]
                    )
                    break

        missing = [
            field_name
            for field_name in aliases
            if field_name not in resolved
        ]

        if missing:
            readable = {
                "code": "Código",
                "name": "Material",
                "quantity": "Quantidade",
                "unit": "Unidade",
            }

            names = ", ".join(
                readable[item]
                for item in missing
            )

            raise ValueError(
                "Colunas obrigatórias não encontradas: "
                f"{names}."
            )

        return resolved

    @staticmethod
    def _normalize_header(value) -> str:
        if value is None:
            return ""

        text = str(value).strip().lower()

        replacements = {
            "á": "a",
            "à": "a",
            "â": "a",
            "ã": "a",
            "ä": "a",
            "é": "e",
            "è": "e",
            "ê": "e",
            "ë": "e",
            "í": "i",
            "ì": "i",
            "î": "i",
            "ï": "i",
            "ó": "o",
            "ò": "o",
            "ô": "o",
            "õ": "o",
            "ö": "o",
            "ú": "u",
            "ù": "u",
            "û": "u",
            "ü": "u",
            "ç": "c",
        }

        for old, new in replacements.items():
            text = text.replace(old, new)

        return " ".join(
            text.split()
        )

    @staticmethod
    def _cell(
        values,
        index: int,
    ):
        if index >= len(values):
            return None

        return values[index]

    @staticmethod
    def _text(value) -> str:
        if value is None:
            return ""

        return str(value).strip()

    @staticmethod
    def _quantity(value) -> float:
        if value is None:
            raise ValueError(
                "Quantidade vazia."
            )

        if isinstance(value, (int, float)):
            return float(value)

        text = str(value).strip()

        if not text:
            raise ValueError(
                "Quantidade vazia."
            )

        # Permite planilhas portuguesas:
        # 1.250,50 -> 1250.50
        if "," in text:
            text = text.replace(".", "")
            text = text.replace(",", ".")

        return float(text)

    @staticmethod
    def _row_is_empty(values) -> bool:
        return all(
            value is None
            or str(value).strip() == ""
            for value in values
        )