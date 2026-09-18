"""Correção Etapa 21 - parser de PDF digital tabular."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path


class QuoteImportError(Exception):
    pass


class QuoteImportService:
    HEADERS = {
        "code": {"codigo", "cod", "referencia", "ref", "reference", "sku", "codigo produto", "codigo material"},
        "description": {"descricao", "designacao", "produto", "material", "artigo", "description", "item"},
        "unit": {"unidade", "un", "unit", "ud", "um"},
        "quantity": {"quantidade", "qtd", "qt", "quantity", "qty"},
        "price": {"preco unitario", "preco", "valor unitario", "valor", "unit price", "price", "p unit", "preco un"},
    }

    def __init__(self, repo, suppliers, inventory):
        self.repo = repo
        self.suppliers = suppliers
        self.inventory = inventory

    def norm(self, value):
        text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
        return re.sub(r"\s+", " ", "".join(c for c in text if not unicodedata.combining(c)))

    def number(self, value):
        if value is None or value == "":
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = re.sub(r"[€$£\s]", "", str(value).strip())
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
        elif "," in text:
            text = text.replace(".", "").replace(",", ".")
        try:
            return float(text)
        except ValueError:
            return None

    def rows(self, name, data):
        ext = Path(name).suffix.lower()
        if ext == ".xlsx":
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).active
            return [list(r) for r in ws.iter_rows(values_only=True)], ext
        if ext == ".csv":
            text = data.decode("utf-8-sig", errors="replace")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\\t")
            except csv.Error:
                dialect = csv.excel
                dialect.delimiter = ";"
            return list(csv.reader(io.StringIO(text), dialect)), ext
        if ext == ".pdf":
            return self._pdf_rows(data), ext
        raise QuoteImportError("Formato não suportado. Use .xlsx, .csv ou PDF digital.")

    def _pdf_rows(self, data):
        try:
            import pdfplumber
        except ImportError as exc:
            raise QuoteImportError("Dependência pdfplumber ausente. Execute: pip install pdfplumber") from exc

        rows = []
        text_fragments = []

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                # Estratégia principal: extrair tabelas preservando células.
                tables = page.extract_tables(
                    table_settings={
                        "vertical_strategy": "lines",
                        "horizontal_strategy": "lines",
                        "intersection_tolerance": 5,
                        "snap_tolerance": 4,
                        "join_tolerance": 4,
                    }
                )
                for table in tables or []:
                    for row in table or []:
                        clean = [re.sub(r"\s+", " ", str(cell or "")).strip() for cell in row]
                        if any(clean):
                            rows.append(clean)

                # Texto fica disponível para fallback.
                text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
                if text.strip():
                    text_fragments.extend(line.strip() for line in text.splitlines() if line.strip())

        if rows:
            return rows

        # Fallback 1: texto visualmente alinhado por espaços.
        fallback_rows = []
        for line in text_fragments:
            cells = [x.strip() for x in re.split(r"\s{2,}", line) if x.strip()]
            if len(cells) >= 2:
                fallback_rows.append(cells)
        if fallback_rows:
            return fallback_rows

        # Fallback 2: devolve linhas simples; pdfitems tentará reconstruir blocos.
        if text_fragments:
            return [[line] for line in text_fragments]

        raise QuoteImportError("PDF sem texto digital extraível. OCR não está habilitado nesta etapa.")

    def tabular(self, rows):
        best = (-1, {})
        for i, row in enumerate(rows[:40]):
            mapping = {}
            for j, cell in enumerate(row):
                normalized = self.norm(cell)
                for field, aliases in self.HEADERS.items():
                    if normalized in aliases and field not in mapping:
                        mapping[field] = j
            score = len(mapping) + (3 if "price" in mapping else 0)
            old_score = len(best[1]) + (3 if "price" in best[1] else 0)
            if score > old_score:
                best = (i, mapping)

        i, mapping = best
        if i < 0 or "price" not in mapping or ("code" not in mapping and "description" not in mapping):
            raise QuoteImportError("Não identifiquei as colunas de preço e material.")

        output = []
        for row_number, row in enumerate(rows[i + 1:], start=i + 2):
            def get(field):
                col = mapping.get(field)
                return row[col] if col is not None and col < len(row) else None

            price = self.number(get("price"))
            ref = str(get("code") or "").strip()
            desc = str(get("description") or "").strip()
            if price is None or price < 0 or not (ref or desc):
                continue
            output.append({
                "row_number": row_number,
                "supplier_reference": ref,
                "description": desc,
                "unit": str(get("unit") or "").strip(),
                "quantity": self.number(get("quantity")),
                "unit_price": price,
                "raw_data": {"ref": ref, "description": desc},
            })
        if not output:
            raise QuoteImportError("Nenhum item válido encontrado.")
        return output

    def pdfitems(self, rows):
        # Primeiro tenta tratar o PDF exatamente como tabela.
        try:
            return self.tabular(rows)
        except QuoteImportError:
            pass

        # Fallback para PDFs cujo extrator devolve uma linha visual completa.
        output = []
        money = re.compile(r"(?:€\s*)?(\d+(?:[.,]\d{1,4}))\s*€?$")
        for rn, row in enumerate(rows, 1):
            line = "  ".join(str(x or "").strip() for x in row if str(x or "").strip())
            match = money.search(line)
            if not match:
                continue
            price = self.number(match.group(1))
            prefix = line[:match.start()].strip(" -|;")
            parts = [p.strip() for p in re.split(r"\s{2,}|[;|]", prefix) if p.strip()]
            ref = parts[0] if len(parts) >= 2 and len(parts[0]) <= 40 else ""
            desc = " ".join(parts[1:]) if ref else prefix
            output.append({
                "row_number": rn,
                "supplier_reference": ref,
                "description": desc,
                "unit": "",
                "quantity": None,
                "unit_price": price,
                "raw_data": {"line": line},
            })
        if not output:
            raise QuoteImportError("PDF legível, mas sem tabela/linhas de preço reconhecíveis.")
        return output

    def detect_supplier(self, name, items, sid):
        if sid is not None:
            return self.suppliers.get_supplier(int(sid))
        hay = self.norm(name + " " + " ".join(x["description"] for x in items[:10]))
        found = []
        for supplier in self.suppliers.list_suppliers(active_only=True):
            if self.norm(supplier.code) in hay or self.norm(supplier.name) in hay:
                found.append(supplier)
        if len(found) == 1:
            return found[0]
        raise QuoteImportError("Fornecedor não identificado com segurança. Selecione-o antes de importar.")

    def match_material(self, sid, item):
        ref = item["supplier_reference"].strip()
        desc = item["description"].strip()
        if ref:
            alias = self.repo.alias(sid, ref)
            if alias:
                return alias, "supplier_alias", 1.0
            try:
                return self.inventory.get_material(ref).code, "internal_code", 1.0
            except Exception:
                pass

        best = (None, 0.0)
        for material in self.inventory.list_materials():
            score = max(
                SequenceMatcher(None, self.norm(desc), self.norm(material.name)).ratio() if desc else 0,
                SequenceMatcher(None, self.norm(ref + " " + desc), self.norm(material.code + " " + material.name)).ratio(),
            )
            if score > best[1]:
                best = (material.code, score)
        return (best[0], "semantic_fuzzy", round(best[1], 4)) if best[1] >= 0.86 else (None, "review", round(best[1], 4))

    def current(self, sid, code):
        try:
            return float(self.suppliers.get_supplier_material(sid, code).unit_price)
        except Exception:
            return None

    def import_file(self, name, data, sid=None):
        rows, ext = self.rows(name, data)
        items = self.pdfitems(rows) if ext == ".pdf" else self.tabular(rows)
        supplier = self.detect_supplier(name, items, sid)
        for item in items:
            code, method, confidence = self.match_material(supplier.id, item)
            item.update(
                matched_material_code=code,
                match_method=method,
                confidence=confidence,
                previous_price=self.current(supplier.id, code) if code else None,
            )
        return self.get(self.repo.create(supplier, name, ext.lstrip("."), items))

    def hydrate(self, record):
        result = dict(record)
        result["items"] = self.repo.items(result["id"])
        for item in result["items"]:
            old, new = item.get("previous_price"), item.get("unit_price")
            item["variation_percent"] = ((new - old) / old * 100) if old not in (None, 0) else None
        return result

    def list(self):
        return [self.hydrate(x) for x in self.repo.list()]

    def get(self, iid):
        record = self.repo.get(iid)
        if not record:
            raise QuoteImportError(f"Importação não encontrada: {iid}.")
        return self.hydrate(record)

    def set_match(self, iid, item_id, code):
        self.get(iid)
        material = self.inventory.get_material(code)
        self.repo.match(item_id, material.code)
        return self.get(iid)

    def approve(self, iid):
        quote = self.get(iid)
        if quote["status"] == "approved":
            raise QuoteImportError("Cotação já aprovada.")
        pending = [x for x in quote["items"] if not x["matched_material_code"]]
        if pending:
            raise QuoteImportError(f"Existem {len(pending)} item(ns) sem material confirmado.")
        sid = int(quote["supplier_id"])
        for item in quote["items"]:
            code = item["matched_material_code"]
            current = self.suppliers.get_supplier_material(sid, code)
            lead_time_days = current.lead_time_days if current else 0
            minimum_order_quantity = current.minimum_order_quantity if current else 0.0
            self.suppliers.upsert_material(sid, code, float(item["unit_price"]), lead_time_days, minimum_order_quantity)
            self.repo.save_alias(sid, item["supplier_reference"], code, item["description"])
        self.repo.approve(iid)
        return self.get(iid)
