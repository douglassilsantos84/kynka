from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import mimetypes
import os
import re
import shutil
import urllib.error
import urllib.request
import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path

from kynka.infrastructure.documents import (
    DocumentDuplicateError,
    DocumentNotFoundError,
)


SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".docx", ".csv", ".txt"}


class DocumentUnsupportedError(Exception):
    pass


class DocumentExtractionError(Exception):
    pass


class EmbeddingEngine:
    """
    Usa Ollama quando disponível. Se não houver modelo de embeddings,
    cai para um vetor lexical determinístico local para manter o RAG funcional.
    """

    def __init__(self):
        self.base_url = os.getenv(
            "KYNKA_OLLAMA_URL",
            "http://127.0.0.1:11434",
        ).rstrip("/")
        self.model = os.getenv(
            "KYNKA_EMBEDDING_MODEL",
            "nomic-embed-text",
        ).strip()
        self.dimension = int(
            os.getenv("KYNKA_LOCAL_EMBEDDING_DIMENSION", "384")
        )
        self._ollama_unavailable = False

    def embed(self, text):
        text = (text or "").strip()
        if not text:
            return [], "empty"

        vector = self._ollama_embed(text)
        if vector:
            return vector, f"ollama:{self.model}"

        return self._local_hash_embedding(text), "local-hash"

    def _ollama_embed(self, text):
        if not self.model or self._ollama_unavailable:
            return None

        requests_to_try = [
            (
                f"{self.base_url}/api/embed",
                {"model": self.model, "input": text},
                "embeddings",
            ),
            (
                f"{self.base_url}/api/embeddings",
                {"model": self.model, "prompt": text},
                "embedding",
            ),
        ]

        for url, payload, key in requests_to_try:
            try:
                req = urllib.request.Request(
                    url,
                    data=json.dumps(payload).encode("utf-8"),
                    headers={"Content-Type": "application/json"},
                    method="POST",
                )
                with urllib.request.urlopen(req, timeout=4) as response:
                    data = json.loads(response.read().decode("utf-8"))
                value = data.get(key)
                if key == "embeddings" and value:
                    value = value[0]
                if isinstance(value, list) and value:
                    return [float(x) for x in value]
            except Exception:
                continue

        self._ollama_unavailable = True
        return None

    def _local_hash_embedding(self, text):
        tokens = re.findall(
            r"[A-Za-zÀ-ÿ0-9_\-]{2,}",
            text.lower(),
        )
        counts = Counter(tokens)
        vector = [0.0] * self.dimension

        for token, count in counts.items():
            digest = hashlib.sha256(token.encode("utf-8")).digest()
            index = int.from_bytes(digest[:4], "big") % self.dimension
            sign = 1.0 if digest[4] % 2 == 0 else -1.0
            weight = 1.0 + math.log(max(1, count))
            vector[index] += sign * weight

        norm = math.sqrt(sum(x * x for x in vector))
        if norm:
            vector = [x / norm for x in vector]
        return vector


class DocumentService:
    def __init__(self, repository, storage_directory):
        self.repository = repository
        self.storage = Path(storage_directory)
        self.storage.mkdir(parents=True, exist_ok=True)
        self.embedding = EmbeddingEngine()

    # ---------------------------------------------------------
    # Public API
    # ---------------------------------------------------------

    def upload(
        self,
        *,
        file_name,
        content,
        mime_type=None,
        title=None,
        category=None,
    ):
        ext = Path(file_name).suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise DocumentUnsupportedError(
                "Formato não suportado. Use PDF, XLSX, DOCX, CSV ou TXT."
            )

        sha256 = hashlib.sha256(content).hexdigest()
        existing = self.repository.get_document_by_sha(sha256)
        if existing:
            raise DocumentDuplicateError(
                f"Este arquivo já existe como documento #{existing['id']}."
            )

        extracted_text = self._extract(file_name, content)
        if not extracted_text.strip():
            raise DocumentExtractionError(
                "Nenhum texto útil foi extraído do documento."
            )

        detected_category = category or self._classify(
            file_name,
            extracted_text,
        )
        metadata = self._extract_metadata(
            file_name,
            extracted_text,
        )

        stored_name = (
            f"{uuid.uuid4().hex}"
            f"{ext}"
        )
        target = self.storage / stored_name
        target.write_bytes(content)

        try:
            document_id = self.repository.create_document(
                original_name=file_name,
                stored_name=stored_name,
                file_type=ext.lstrip("."),
                mime_type=mime_type
                or mimetypes.guess_type(file_name)[0]
                or "application/octet-stream",
                size_bytes=len(content),
                sha256=sha256,
                title=(title or Path(file_name).stem).strip(),
                category=detected_category,
                extracted_text=extracted_text,
                metadata=metadata,
            )

            chunks = self._build_chunks(extracted_text)
            self.repository.replace_chunks(
                document_id,
                chunks,
            )
        except Exception:
            target.unlink(missing_ok=True)
            raise

        return self.get(document_id)

    def list(self, category=None, search=None, limit=100):
        return self.repository.list_documents(
            category=category,
            search=search,
            limit=limit,
        )

    def get(self, document_id):
        document = self.repository.get_document(document_id)
        document["chunks"] = self.repository.get_chunks(document_id)
        for chunk in document["chunks"]:
            chunk.pop("embedding", None)
        return document

    def delete(self, document_id):
        document = self.repository.get_document(document_id)
        self.repository.delete_document(document_id)
        stored = document.get("stored_name")
        if stored:
            (self.storage / stored).unlink(missing_ok=True)
        return {"success": True, "document_id": int(document_id)}

    def reindex(self, document_id):
        document = self.repository.get_document(document_id)
        stored = self.storage / document["stored_name"]
        if not stored.exists():
            raise DocumentExtractionError(
                "Arquivo físico do documento não foi encontrado."
            )

        content = stored.read_bytes()
        extracted_text = self._extract(
            document["original_name"],
            content,
        )
        category = self._classify(
            document["original_name"],
            extracted_text,
        )
        metadata = self._extract_metadata(
            document["original_name"],
            extracted_text,
        )
        self.repository.update_index(
            document_id,
            extracted_text=extracted_text,
            category=category,
            metadata=metadata,
        )
        self.repository.replace_chunks(
            document_id,
            self._build_chunks(extracted_text),
        )
        return self.get(document_id)

    def stats(self):
        result = self.repository.stats()
        result["supported_formats"] = sorted(
            x.lstrip(".") for x in SUPPORTED_EXTENSIONS
        )
        result["embedding_model"] = self.embedding.model
        return result

    def search(self, query, document_ids=None, limit=6):
        query = (query or "").strip()
        if not query:
            return []

        query_vector, query_provider = self.embedding.embed(query)
        query_terms = set(self._terms(query))
        candidates = self.repository.all_chunks(document_ids=document_ids)

        scored = []
        for item in candidates:
            content = item["content"]
            vector = item["embedding"]

            semantic = self._cosine(query_vector, vector)
            content_terms = set(self._terms(content))
            lexical = (
                len(query_terms & content_terms)
                / max(1, len(query_terms))
            )

            # Prioriza semântica quando o vetor tem dimensão compatível,
            # mas mantém recuperação lexical robusta como fallback.
            if (
                query_vector
                and vector
                and len(query_vector) == len(vector)
            ):
                score = (semantic * 0.78) + (lexical * 0.22)
            else:
                score = lexical

            if score <= 0:
                continue

            scored.append(
                {
                    "score": round(float(score), 6),
                    "document_id": item["document_id"],
                    "document_name": item["original_name"],
                    "title": item["title"],
                    "category": item["category"],
                    "chunk_index": item["chunk_index"],
                    "content": content,
                    "embedding_provider": item["embedding_provider"],
                    "query_embedding_provider": query_provider,
                }
            )

        scored.sort(
            key=lambda x: x["score"],
            reverse=True,
        )
        return self._diversify_sources(
            query,
            scored,
            max(1, int(limit)),
        )

    def ask(self, question, document_ids=None, limit=6):
        sources = self.search(
            question,
            document_ids=document_ids,
            limit=limit,
        )

        if not sources:
            return {
                "answer": (
                    "Não encontrei informação suficiente nos documentos "
                    "indexados para responder com segurança."
                ),
                "sources": [],
                "grounded": True,
                "mode": "no_context",
            }

        grounded_answer = self._grounded_material_answer(
            question,
            sources,
        )

        if grounded_answer:
            answer = grounded_answer
            mode = "grounded-structured-rag"
        else:
            answer = self._answer_with_ollama(
                question,
                sources,
            )
            mode = "ollama-rag"

            if not answer:
                answer = self._extractive_answer(
                    question,
                    sources,
                )
                mode = "extractive-rag"

        clean_sources = []
        for source in sources:
            clean_sources.append(
                {
                    "document_id": source["document_id"],
                    "document_name": source["document_name"],
                    "title": source["title"],
                    "category": source["category"],
                    "chunk_index": source["chunk_index"],
                    "score": source["score"],
                    "excerpt": self._excerpt(
                        source["content"],
                        420,
                    ),
                }
            )

        return {
            "answer": answer,
            "sources": clean_sources,
            "grounded": True,
            "mode": mode,
        }

    def _grounded_material_answer(self, question, sources):
        """
        Para perguntas compostas sobre materiais, extrai fatos diretamente
        das evidências antes de envolver o LLM. Isso impede conversão de moeda,
        confusão entre preço unitário e total e omissão de campos explícitos.
        """
        q = (question or "").lower()
        requested = {
            "description": any(x in q for x in ("descrição", "descricao")),
            "quantity": "quantidade" in q,
            "price": "preço" in q or "preco" in q,
            "supplier": "fornecedor" in q,
            "delivery": "prazo de entrega" in q or "entrega" in q,
            "total": "valor total" in q or "preço total" in q or "preco total" in q,
            "reference": "referência" in q or "referencia" in q,
        }
        if sum(requested.values()) < 2:
            return None

        reference_match = re.search(
            r"\b[A-Z]{2,}[A-Z0-9]*[-_/][A-Z0-9][A-Z0-9\-_/]*\b",
            question or "",
            re.I,
        )
        target_reference = reference_match.group(0) if reference_match else None

        facts = {
            "description": None,
            "quantity": None,
            "price": None,
            "supplier": None,
            "delivery": None,
            "total": None,
            "reference": None,
        }

        for source_index, source in enumerate(sources, start=1):
            content = source["content"]
            normalized = re.sub(r"\s+", " ", content).strip()

            if target_reference and target_reference.lower() not in normalized.lower():
                continue

            # Tabelas CSV/XLSX: Referencia | Descricao | Quantidade | Preco
            table = re.search(
                r"(?:Referencia|Referência)\s*\|\s*Descricao\s*\|\s*Quantidade\s*\|\s*Preco"
                r".*?\b([A-Z]{2,}[A-Z0-9]*[-_/][A-Z0-9][A-Z0-9\-_/]*)\s*\|\s*"
                r"([^|\n]+?)\s*\|\s*([0-9]+(?:[.,][0-9]+)?)\s*\|\s*"
                r"([0-9]+(?:[.,][0-9]+)?)(?:\s*([€$]|EUR|USD|R\$))?",
                normalized,
                re.I,
            )
            if table:
                ref, desc, qty, price, currency = table.groups()
                if not target_reference or ref.lower() == target_reference.lower():
                    facts["reference"] = facts["reference"] or (ref, source_index)
                    facts["description"] = facts["description"] or (desc.strip(), source_index)
                    facts["quantity"] = facts["quantity"] or (qty.strip(), source_index)
                    price_value = price.strip()
                    if currency:
                        price_value = f"{price_value} {currency.strip()}"
                    else:
                        price_value = f"{price_value} (moeda não informada)"
                    facts["price"] = facts["price"] or (price_value, source_index)

            supplier = re.search(
                r"\bFornecedor\s*:\s*([^\n\r|]+?)(?=\s+NIF\s*:|\s+Prazo\s+de\s+entrega\s*:|$)",
                normalized,
                re.I,
            )
            if supplier:
                facts["supplier"] = facts["supplier"] or (
                    supplier.group(1).strip(),
                    source_index,
                )

            delivery = re.search(
                r"\bPrazo\s+de\s+entrega\s*:\s*([^\n\r|.;]+)",
                normalized,
                re.I,
            )
            if delivery:
                facts["delivery"] = facts["delivery"] or (
                    delivery.group(1).strip(),
                    source_index,
                )

            total = re.search(
                r"\bValor\s+total\s*:\s*"
                r"([0-9]{1,3}(?:[.\s][0-9]{3})*(?:,[0-9]{2})?)\s*"
                r"(€|EUR|R\$|USD|\$)?",
                normalized,
                re.I,
            )
            if total:
                value = total.group(1).strip()
                currency = (total.group(2) or "").strip()
                total_value = f"{value} {currency}".strip()
                if not currency:
                    total_value += " (moeda não informada)"
                facts["total"] = facts["total"] or (
                    total_value,
                    source_index,
                )

            if target_reference and target_reference.lower() in normalized.lower():
                facts["reference"] = facts["reference"] or (
                    target_reference,
                    source_index,
                )

        labels = [
            ("description", "Descrição"),
            ("quantity", "Quantidade"),
            ("price", "Preço"),
            ("supplier", "Fornecedor"),
            ("delivery", "Prazo de entrega"),
            ("total", "Valor total"),
            ("reference", "Referência"),
        ]

        lines = []
        for key, label in labels:
            if not requested[key]:
                continue
            fact = facts[key]
            if fact:
                value, source_index = fact
                lines.append(f"**{label}:** {value} [Fonte {source_index}]")
            else:
                lines.append(f"**{label}:** não encontrado nas fontes")

        return "\n\n".join(lines) if lines else None

    # ---------------------------------------------------------
    # Parsing
    # ---------------------------------------------------------

    def _extract(self, file_name, content):
        ext = Path(file_name).suffix.lower()
        try:
            if ext == ".pdf":
                return self._extract_pdf(content)
            if ext == ".xlsx":
                return self._extract_xlsx(content)
            if ext == ".docx":
                return self._extract_docx(content)
            if ext == ".csv":
                return self._extract_csv(content)
            if ext == ".txt":
                return self._decode_text(content)
        except DocumentExtractionError:
            raise
        except Exception as exc:
            raise DocumentExtractionError(
                f"Falha ao extrair {file_name}: {exc}"
            ) from exc
        raise DocumentUnsupportedError(
            f"Formato não suportado: {ext}"
        )

    def _extract_pdf(self, content):
        try:
            import pdfplumber
        except ImportError as exc:
            raise DocumentExtractionError(
                "Dependência pdfplumber não instalada."
            ) from exc

        parts = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                tables = page.extract_tables() or []
                table_text = []
                for table in tables:
                    for row in table:
                        values = [
                            str(cell).strip()
                            if cell is not None
                            else ""
                            for cell in row
                        ]
                        table_text.append(" | ".join(values))
                page_content = "\n".join(
                    x for x in [
                        text.strip(),
                        "\n".join(table_text).strip(),
                    ]
                    if x
                )
                if page_content:
                    parts.append(
                        f"[Página {index}]\n{page_content}"
                    )

        return "\n\n".join(parts)

    def _extract_xlsx(self, content):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise DocumentExtractionError(
                "Dependência openpyxl não instalada."
            ) from exc

        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
        parts = []
        for sheet in workbook.worksheets:
            rows = [f"[Planilha: {sheet.title}]"]
            for row in sheet.iter_rows(values_only=True):
                values = [
                    "" if value is None else str(value).strip()
                    for value in row
                ]
                if any(values):
                    rows.append(" | ".join(values))
            parts.append("\n".join(rows))
        return "\n\n".join(parts)

    def _extract_docx(self, content):
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentExtractionError(
                "Dependência python-docx não instalada."
            ) from exc

        document = Document(io.BytesIO(content))
        parts = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table_index, table in enumerate(document.tables, start=1):
            parts.append(f"[Tabela {table_index}]")
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                ]
                if any(values):
                    parts.append(" | ".join(values))

        return "\n".join(parts)

    def _extract_csv(self, content):
        text = self._decode_text(content)
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(
                sample,
                delimiters=";,|\t,",
            )
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"

        reader = csv.reader(io.StringIO(text), dialect)
        rows = []
        for row in reader:
            values = [cell.strip() for cell in row]
            if any(values):
                rows.append(" | ".join(values))
        return "\n".join(rows)

    @staticmethod
    def _decode_text(content):
        for encoding in (
            "utf-8-sig",
            "utf-8",
            "cp1252",
            "latin-1",
        ):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    # ---------------------------------------------------------
    # Intelligence
    # ---------------------------------------------------------

    def _classify(self, file_name, text):
        hay = f"{file_name}\n{text[:9000]}".lower()

        rules = [
            (
                "cotacao",
                (
                    "cotação",
                    "cotacao",
                    "orçamento",
                    "orcamento",
                    "proposta comercial",
                    "quotation",
                    "quote",
                ),
            ),
            (
                "fatura",
                (
                    "fatura",
                    "factura",
                    "invoice",
                    "n.º fatura",
                    "nº fatura",
                ),
            ),
            (
                "contrato",
                (
                    "contrato",
                    "contract",
                    "cláusula",
                    "clausula",
                    "outorgante",
                ),
            ),
            (
                "projeto",
                (
                    "mapa de quantidades",
                    "memória descritiva",
                    "memoria descritiva",
                    "projeto",
                    "obra",
                ),
            ),
            (
                "manual",
                (
                    "manual",
                    "instruções",
                    "instrucoes",
                    "datasheet",
                    "ficha técnica",
                    "ficha tecnica",
                ),
            ),
        ]

        for category, words in rules:
            if any(word in hay for word in words):
                return category
        return "geral"

    def _extract_metadata(self, file_name, text):
        sample = text[:50000]

        nifs = sorted(
            set(
                re.findall(
                    r"(?<!\d)\d{9}(?!\d)",
                    sample,
                )
            )
        )

        emails = sorted(
            set(
                re.findall(
                    r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b",
                    sample,
                )
            )
        )

        dates = sorted(
            set(
                re.findall(
                    r"\b(?:0?[1-9]|[12]\d|3[01])[\/\-.]"
                    r"(?:0?[1-9]|1[0-2])[\/\-.](?:19|20)\d{2}\b",
                    sample,
                )
            )
        )[:25]

        euro_values = sorted(
            set(
                re.findall(
                    r"(?:€\s*)?\d{1,3}(?:[.\s]\d{3})*(?:,\d{2})\s*€?",
                    sample,
                )
            )
        )[:40]

        references = sorted(
            set(
                re.findall(
                    r"\b[A-Z]{2,}[A-Z0-9]*[-_/][A-Z0-9][A-Z0-9\-_/]*\b",
                    sample,
                )
            )
        )[:80]

        supplier = self._first_metadata_match(
            sample,
            [
                r"(?im)^\s*(?:fornecedor|supplier)\s*[:\-]\s*([^\n\r|]{2,120})",
                r"(?im)^\s*(?:empresa|company)\s*[:\-]\s*([^\n\r|]{2,120})",
            ],
        )

        delivery_time = self._first_metadata_match(
            sample,
            [
                r"(?im)\b(?:prazo\s+de\s+entrega|delivery\s+(?:time|term))\s*[:\-]?\s*([^\n\r|.;]{1,80})",
                r"(?im)\bentrega\s+(?:em|dentro\s+de)\s+([^\n\r|.;]{1,80})",
            ],
        )

        proposal_validity = self._first_metadata_match(
            sample,
            [
                r"(?im)\b(?:validade(?:\s+da\s+(?:proposta|cotação|cotacao))?|proposal\s+validity|quote\s+validity)\s*[:\-]?\s*([^\n\r|.;]{1,100})",
                r"(?im)\b(?:proposta|cotação|cotacao)\s+válida\s+(?:por|até)\s+([^\n\r|.;]{1,100})",
            ],
        )

        payment_terms = self._first_metadata_match(
            sample,
            [
                r"(?im)\b(?:condiç(?:ão|ões)\s+de\s+pagamento|condicoes\s+de\s+pagamento|payment\s+terms?)\s*[:\-]?\s*([^\n\r|.;]{1,160})",
                r"(?im)\bpagamento\s*[:\-]\s*([^\n\r|.;]{1,160})",
            ],
        )

        document_date = self._first_metadata_match(
            sample,
            [
                r"(?im)\b(?:data(?:\s+do\s+documento)?|date)\s*[:\-]\s*((?:0?[1-9]|[12]\d|3[01])[\/\-.](?:0?[1-9]|1[0-2])[\/\-.](?:19|20)\d{2})",
            ],
        )

        return {
            "file_name": file_name,
            "supplier": supplier,
            "nifs": nifs[:20],
            "emails": emails[:20],
            "document_date": document_date,
            "dates": dates,
            "delivery_time": delivery_time,
            "proposal_validity": proposal_validity,
            "payment_terms": payment_terms,
            "euro_values": euro_values,
            "references": references,
            "text_characters": len(text),
        }

    @staticmethod
    def _first_metadata_match(text, patterns):
        for pattern in patterns:
            match = re.search(pattern, text)
            if not match:
                continue

            value = re.sub(
                r"\s+",
                " ",
                match.group(1),
            ).strip(" \t:-")

            if value:
                return value

        return None

    def _build_chunks(
        self,
        text,
        chunk_size=1400,
        overlap=220,
    ):
        cleaned = re.sub(
            r"\n{3,}",
            "\n\n",
            text,
        ).strip()

        chunks_text = []
        start = 0
        while start < len(cleaned):
            end = min(
                len(cleaned),
                start + chunk_size,
            )
            if end < len(cleaned):
                boundary = max(
                    cleaned.rfind("\n\n", start, end),
                    cleaned.rfind(". ", start, end),
                    cleaned.rfind("\n", start, end),
                )
                if boundary > start + 500:
                    end = boundary + 1

            chunk = cleaned[start:end].strip()
            if chunk:
                chunks_text.append(chunk)

            if end >= len(cleaned):
                break
            start = max(
                start + 1,
                end - overlap,
            )

        result = []
        for index, chunk in enumerate(chunks_text):
            vector, provider = self.embedding.embed(chunk)
            result.append(
                {
                    "chunk_index": index,
                    "content": chunk,
                    "embedding": vector,
                    "embedding_provider": provider,
                }
            )
        return result

    def _answer_with_ollama(self, question, sources):
        model = os.getenv(
            "KYNKA_DOCUMENT_CHAT_MODEL",
            os.getenv("KYNKA_MODEL", "llama3.2:3b"),
        ).strip()
        if not model:
            return None

        context_parts = []
        for idx, source in enumerate(sources, start=1):
            context_parts.append(
                f"[Fonte {idx} | documento #{source['document_id']} | "
                f"{source['document_name']} | trecho {source['chunk_index']}]\n"
                f"{source['content']}"
            )

        system = (
            "Você é a Kynka. Responda exclusivamente com base nas fontes "
            "fornecidas. Não invente dados. Se as fontes forem insuficientes, "
            "diga claramente que não há informação suficiente. "
            "Ao usar uma informação, indique [Fonte N]."
        )
        user = (
            f"Pergunta: {question}\n\n"
            "FONTES:\n"
            + "\n\n".join(context_parts)
        )

        payload = {
            "model": model,
            "stream": False,
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            "options": {
                "temperature": 0.1,
            },
        }

        try:
            url = (
                os.getenv(
                    "KYNKA_OLLAMA_URL",
                    "http://127.0.0.1:11434",
                ).rstrip("/")
                + "/api/chat"
            )
            request = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json"},
                method="POST",
            )
            with urllib.request.urlopen(
                request,
                timeout=90,
            ) as response:
                data = json.loads(
                    response.read().decode("utf-8")
                )
            message = data.get("message") or {}
            content = (message.get("content") or "").strip()
            return content or None
        except Exception:
            return None

    def _diversify_sources(self, query, scored, limit):
        if not scored:
            return []

        query_terms = set(self._terms(query))
        selected = []
        remaining = list(scored)
        covered_terms = set()

        while remaining and len(selected) < limit:
            best_index = 0
            best_value = float("-inf")

            for index, item in enumerate(remaining):
                content_terms = set(self._terms(item["content"]))
                new_terms = (query_terms & content_terms) - covered_terms
                novelty = len(new_terms) / max(1, len(query_terms))

                normalized = set(self._terms(item["content"]))
                duplicate = False
                for chosen in selected:
                    other = set(self._terms(chosen["content"]))
                    union = normalized | other
                    similarity = (
                        len(normalized & other) / len(union)
                        if union else 0.0
                    )
                    if similarity >= 0.88:
                        duplicate = True
                        break

                value = (
                    float(item["score"])
                    + novelty * 0.65
                    - (0.22 if duplicate else 0.0)
                )
                if value > best_value:
                    best_value = value
                    best_index = index

            chosen = remaining.pop(best_index)
            selected.append(chosen)
            covered_terms.update(
                query_terms & set(self._terms(chosen["content"]))
            )

        return selected

    def _extractive_answer(self, question, sources):
        terms = set(self._terms(question))
        sentences = []

        for source_index, source in enumerate(sources, start=1):
            raw_sentences = re.split(
                r"(?<=[.!?])\s+|\n+",
                source["content"],
            )
            for sentence in raw_sentences:
                sentence = sentence.strip()
                if len(sentence) < 20:
                    continue
                sentence_terms = set(
                    self._terms(sentence)
                )
                overlap = len(
                    terms & sentence_terms
                )
                score = (
                    overlap
                    + float(source["score"]) * 2
                )
                sentences.append(
                    (
                        score,
                        sentence,
                        source_index,
                    )
                )

        sentences.sort(
            key=lambda item: item[0],
            reverse=True,
        )

        selected = []
        seen = set()
        for _, sentence, source_index in sentences:
            key = sentence.lower()
            if key in seen:
                continue
            seen.add(key)
            selected.append(
                f"{sentence} [Fonte {source_index}]"
            )
            if len(selected) >= 4:
                break

        if not selected:
            return (
                "Encontrei documentos relacionados, mas não há trecho "
                "suficientemente claro para responder sem inferir."
            )

        return " ".join(selected)

    @staticmethod
    def _terms(text):
        stop = {
            "a", "ao", "aos", "as", "o", "os", "de", "da", "das",
            "do", "dos", "e", "em", "um", "uma", "para", "por",
            "com", "que", "qual", "quais", "como", "no", "na",
            "nos", "nas", "é", "são", "se", "me", "meu", "minha",
        }
        return [
            token
            for token in re.findall(
                r"[A-Za-zÀ-ÿ0-9_\-]{2,}",
                (text or "").lower(),
            )
            if token not in stop
        ]

    @staticmethod
    def _cosine(a, b):
        if (
            not a
            or not b
            or len(a) != len(b)
        ):
            return 0.0
        dot = sum(x * y for x, y in zip(a, b))
        na = math.sqrt(sum(x * x for x in a))
        nb = math.sqrt(sum(y * y for y in b))
        if not na or not nb:
            return 0.0
        return dot / (na * nb)

    @staticmethod
    def _excerpt(text, limit):
        text = re.sub(r"\s+", " ", text).strip()
        if len(text) <= limit:
            return text
        return text[: limit - 1].rstrip() + "…"
