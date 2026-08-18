"""
Importação de inventário para a plataforma Kynka.

Formatos suportados:
- CSV
- XLSX
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from kynka.application.inventory import InventoryService


class InventoryImportError(Exception):
    """
    Erro durante a importação de inventário.
    """


@dataclass(slots=True)
class InventoryImportResult:
    """
    Resultado de uma importação.
    """

    imported: int = 0
    skipped: int = 0
    errors: list[str] | None = None

    def __post_init__(self) -> None:
        if self.errors is None:
            self.errors = []


class InventoryImporter:
    """
    Importa materiais de arquivos CSV ou XLSX.
    """

    COLUMN_ALIASES = {
        "code": {
            "code",
            "codigo",
            "código",
            "cod",
            "sku",
            "referencia",
            "referência",
            "ref",
        },
        "name": {
            "name",
            "nome",
            "descricao",
            "descrição",
            "material",
            "produto",
        },
        "quantity": {
            "quantity",
            "quantidade",
            "qtd",
            "saldo",
            "estoque",
            "stock",
        },
        "unit": {
            "unit",
            "unidade",
            "un",
        },
        "minimum_quantity": {
            "minimum_quantity",
            "minimum",
            "minimo",
            "mínimo",
            "estoque_minimo",
            "estoque mínimo",
            "stock_minimo",
            "stock mínimo",
            "qtd_minima",
            "qtd mínima",
        },
    }

    def __init__(
        self,
        service: InventoryService,
    ) -> None:
        self._service = service

    def import_file(
        self,
        file_path: str | Path,
    ) -> InventoryImportResult:
        """
        Detecta o formato e importa o arquivo.
        """

        path = Path(file_path)

        if not path.exists():
            raise InventoryImportError(
                f"Arquivo não encontrado: {path}"
            )

        extension = path.suffix.lower()

        if extension == ".csv":
            rows = self._read_csv(path)

        elif extension == ".xlsx":
            rows = self._read_xlsx(path)

        else:
            raise InventoryImportError(
                "Formato não suportado. "
                "Utilize arquivos .csv ou .xlsx."
            )

        return self._import_rows(rows)

    def _read_csv(
        self,
        path: Path,
    ) -> list[dict[str, Any]]:

        encodings = (
            "utf-8-sig",
            "utf-8",
            "cp1252",
            "latin-1",
        )

        last_error: Exception | None = None

        for encoding in encodings:
            try:
                with path.open(
                    "r",
                    encoding=encoding,
                    newline="",
                ) as file:
                    sample = file.read(4096)
                    file.seek(0)

                    try:
                        dialect = csv.Sniffer().sniff(
                            sample,
                            delimiters=",;",
                        )
                    except csv.Error:
                        dialect = csv.excel

                    reader = csv.DictReader(
                        file,
                        dialect=dialect,
                    )

                    return [
                        dict(row)
                        for row in reader
                    ]

            except UnicodeDecodeError as error:
                last_error = error

        raise InventoryImportError(
            "Não foi possível identificar "
            "a codificação do arquivo CSV."
        ) from last_error

    @staticmethod
    def _read_xlsx(
        path: Path,
    ) -> list[dict[str, Any]]:

        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        try:
            worksheet = workbook.active

            iterator = worksheet.iter_rows(
                values_only=True
            )

            try:
                headers = next(iterator)
            except StopIteration:
                return []

            normalized_headers = [
                str(value).strip()
                if value is not None
                else ""
                for value in headers
            ]

            rows: list[dict[str, Any]] = []

            for values in iterator:
                row = {
                    normalized_headers[index]: value
                    for index, value in enumerate(values)
                    if (
                        index < len(normalized_headers)
                        and normalized_headers[index]
                    )
                }

                rows.append(row)

            return rows

        finally:
            workbook.close()

    def _import_rows(
        self,
        rows: list[dict[str, Any]],
    ) -> InventoryImportResult:

        result = InventoryImportResult()

        for index, row in enumerate(
            rows,
            start=2,
        ):
            try:
                normalized = (
                    self._normalize_row(row)
                )

                if normalized is None:
                    result.skipped += 1
                    continue

                self._service.save_material(
                    code=normalized["code"],
                    name=normalized["name"],
                    quantity=normalized["quantity"],
                    unit=normalized["unit"],
                    minimum_quantity=(
                        normalized[
                            "minimum_quantity"
                        ]
                    ),
                )

                result.imported += 1

            except Exception as error:
                result.skipped += 1
                result.errors.append(
                    f"Linha {index}: {error}"
                )

        return result

    def _normalize_row(
        self,
        row: dict[str, Any],
    ) -> dict[str, Any] | None:

        normalized_keys = {
            self._normalize_header(key): value
            for key, value in row.items()
            if key is not None
        }

        code = self._find_value(
            normalized_keys,
            "code",
        )

        name = self._find_value(
            normalized_keys,
            "name",
        )

        quantity = self._find_value(
            normalized_keys,
            "quantity",
        )

        unit = self._find_value(
            normalized_keys,
            "unit",
        )

        minimum = self._find_value(
            normalized_keys,
            "minimum_quantity",
        )

        if self._is_empty(code) and self._is_empty(name):
            return None

        if self._is_empty(code):
            raise InventoryImportError(
                "Código do material ausente."
            )

        if self._is_empty(name):
            raise InventoryImportError(
                "Nome do material ausente."
            )

        return {
            "code": str(code).strip(),
            "name": str(name).strip(),
            "quantity": self._to_float(
                quantity,
                default=0.0,
            ),
            "unit": (
                str(unit).strip()
                if not self._is_empty(unit)
                else "un"
            ),
            "minimum_quantity": self._to_float(
                minimum,
                default=0.0,
            ),
        }

    def _find_value(
        self,
        row: dict[str, Any],
        field: str,
    ) -> Any:

        aliases = {
            self._normalize_header(alias)
            for alias in self.COLUMN_ALIASES[field]
        }

        for key, value in row.items():
            if key in aliases:
                return value

        return None

    @staticmethod
    def _normalize_header(
        value: str,
    ) -> str:

        import unicodedata

        normalized = unicodedata.normalize(
            "NFKD",
            str(value).strip().lower(),
        )

        text = "".join(
            character
            for character in normalized
            if not unicodedata.combining(character)
        )

        return (
            text
            .replace("-", "_")
            .replace(" ", "_")
        )

    @staticmethod
    def _is_empty(
        value: Any,
    ) -> bool:

        return (
            value is None
            or str(value).strip() == ""
        )

    @staticmethod
    def _to_float(
        value: Any,
        default: float = 0.0,
    ) -> float:

        if InventoryImporter._is_empty(value):
            return default

        if isinstance(
            value,
            (int, float),
        ):
            return float(value)

        text = str(value).strip()

        # Formato europeu/brasileiro:
        # 1.234,56 -> 1234.56
        if "," in text:
            text = (
                text
                .replace(".", "")
                .replace(",", ".")
            )

        return float(text)